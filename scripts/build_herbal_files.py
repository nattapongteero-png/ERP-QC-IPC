# -*- coding: utf-8 -*-
"""สร้างไฟล์ผลลัพธ์ 4 ไฟล์ จาก data model + diagrams"""
import os, json, base64, glob, sys

# locate latest data dir
tmp = os.environ.get("TEMP") or os.environ.get("TMP")
cands = sorted(glob.glob(os.path.join(tmp,"herbal_img_*")), key=os.path.getmtime)
IMG = cands[-1]
with open(os.path.join(IMG,"_data.json"),encoding="utf-8") as fp:
    DATA = json.load(fp)
MODULES, FIXES, PERIOD, IMGS = DATA["modules"],DATA["fixes"],DATA["period"],DATA["imgs"]
OUT = r"C:\Users\kook_\OneDrive\Desktop\HerbalERP-คู่มือระบบ\08062569"
TITLE = "Herbal ERP — สรุประบบที่พัฒนา (สัปดาห์ " + PERIOD + ")"

def b64(path):
    with open(path,"rb") as fp:
        return base64.b64encode(fp.read()).decode()

results = []
def saved(name, path):
    sz = os.path.getsize(path)
    results.append((name, path, sz))

# ====================================================================
# 1) HTML  — system-flow_HerbalERP.html
# ====================================================================
def build_html():
    def img_tag(key, w="100%"):
        return f'<img src="data:image/png;base64,{b64(IMGS[key])}" style="max-width:{w};border:1px solid #d8e2ee;border-radius:10px;margin:10px 0;"/>'
    css = """
    *{box-sizing:border-box} body{font-family:'Sarabun','Segoe UI',Tahoma,sans-serif;margin:0;background:#f4f7fb;color:#1f2937;line-height:1.6}
    .wrap{max-width:1180px;margin:0 auto;padding:0 22px 70px}
    header{background:linear-gradient(135deg,#1e3a5f,#2c6fbf);color:#fff;padding:42px 22px;text-align:center;margin-bottom:26px}
    header h1{margin:0 0 8px;font-size:30px} header p{margin:0;opacity:.9}
    h2{color:#1e3a5f;border-left:6px solid #2c6fbf;padding-left:12px;margin-top:42px;font-size:24px}
    h3{color:#0f766e;margin-bottom:4px;font-size:20px}
    .card{background:#fff;border:1px solid #e2e8f0;border-radius:14px;padding:20px 24px;margin:16px 0;box-shadow:0 1px 4px rgba(0,0,0,.05)}
    .menu{display:inline-block;background:#eaf2fb;color:#1e3a5f;border-radius:8px;padding:4px 12px;font-size:14px;margin:6px 0}
    .tag{display:inline-block;border-radius:6px;padding:2px 10px;font-size:13px;font-weight:600;color:#fff;margin-right:6px}
    .t-add{background:#15803d}.t-edit{background:#b45309}.t-del{background:#b91c1c}.t-view{background:#475569}
    table{border-collapse:collapse;width:100%;margin-top:10px;font-size:14px}
    th,td{border:1px solid #dbe3ee;padding:8px 10px;text-align:left;vertical-align:top}
    th{background:#1e3a5f;color:#fff} tr:nth-child(even){background:#f7faff}
    ul{margin:6px 0 6px 0;padding-left:22px} li{margin:3px 0}
    .toc a{display:block;color:#2c6fbf;text-decoration:none;padding:3px 0}
    .legend{font-size:14px;color:#475569}
    .pill{background:#0f766e;color:#fff;border-radius:20px;padding:3px 14px;font-size:13px;display:inline-block;margin:2px}
    .fix h3{color:#b45309}
    footer{text-align:center;color:#94a3b8;font-size:13px;margin-top:50px}
    """
    def tagcls(k):
        return {"เพิ่ม":"t-add","แก้ไข":"t-edit","ลบ":"t-del","ดู":"t-view"}.get(k,"t-view")

    h = [f"<!doctype html><html lang='th'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>",
         "<link href='https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap' rel='stylesheet'>",
         f"<title>{TITLE}</title><style>{css}</style></head><body>"]
    h.append(f"<header><h1>System Flow — Herbal ERP</h1><p>แผนผังการทำงานของระบบ • สัปดาห์ {PERIOD}</p></header><div class='wrap'>")

    # overview
    h.append("<h2>ภาพรวมระบบ</h2><div class='card'>")
    h.append(img_tag("overview"))
    h.append("<p>Herbal ERP จัดการกระบวนการผลิตยาสมุนไพรตามมาตรฐาน GMP ตั้งแต่การตรวจรับวัตถุดิบ การเบิก การชั่ง การผลิต การตรวจคุณภาพ ไปจนถึงการบรรจุและปล่อยสินค้าสำเร็จ โดยมีระบบสนับสนุนด้านสถานที่ ข้อมูลหลัก และเอกสาร GMP กำกับตลอดสาย</p>")
    h.append("<p class='legend'>หลักการสำคัญ: <span class='pill'>QC อิสระจากผู้ผลิต</span><span class='pill'>แยกหน้าที่ (Segregation of Duties)</span><span class='pill'>สอบย้อนกลับได้ (Traceability)</span></p></div>")

    # TOC
    h.append("<h2>สารบัญ</h2><div class='card toc'>")
    h.append("<b>ส่วนที่ 1 — ระบบที่พัฒนาเพิ่ม (เพิ่ม / แก้ไข / ลบ)</b>")
    for i,m in enumerate(MODULES,1):
        h.append(f"<a href='#m{i}'>{i}. {m['name']} <span class='menu'>{m['code']}</span></a>")
    h.append("<br><b>ส่วนที่ 2 — ระบบที่ปรับแก้ไข</b><a href='#fixes'>รายการการปรับแก้ไขระบบทั้งหมด</a>")
    h.append("</div>")

    # PART 1
    h.append("<h2>ส่วนที่ 1 — ระบบที่พัฒนาเพิ่ม</h2>")
    for i,m in enumerate(MODULES,1):
        h.append(f"<div class='card' id='m{i}'><h3>{i}. {m['name']} &nbsp;<span class='menu'>{m['code']}</span></h3>")
        h.append("<table style='width:auto;margin:6px 0'>")
        h.append(f"<tr><th>Module (เมนูหลัก)</th><td><b>{m['module']}</b></td></tr>")
        h.append(f"<tr><th>เมนูที่คลิก</th><td>เมนูซ้าย → {m['module'].split('(')[0].strip()} → <b>{m['menu']}</b></td></tr>")
        h.append(f"<tr><th>URL อ้างอิง</th><td><code>{m['url']}</code></td></tr>")
        if m.get("nav2"):
            h.append(f"<tr><th>เพิ่มเติม</th><td>{m['nav2']}</td></tr>")
        h.append("</table>")
        h.append(f"<p>{m['purpose']}</p>")
        if m.get("diagram"):
            h.append(img_tag(m["diagram"], "760px"))
        for label,key,cls in [("เพิ่ม (Create)","add","t-add"),("แก้ไข (Update)","edit","t-edit"),("ลบ (Delete)","delete","t-del")]:
            if m.get(key):
                h.append(f"<p><span class='tag {cls}'>{label}</span></p><ul>")
                for s in m[key]: h.append(f"<li>{s}</li>")
                h.append("</ul>")
        h.append("</div>")

    # PART 2 fixes — grouped by category
    h.append("<h2 id='fixes'>ส่วนที่ 2 — ระบบที่ปรับแก้ไข</h2><div class='card fix'>")
    h.append("<table><tr><th style='width:18%'>โมดูล</th><th>สิ่งที่ปรับแก้ไข</th><th>วิธีตรวจสอบ (คลิกตามนี้)</th><th style='width:22%'>ผลที่คาดหวัง</th></tr>")
    cur=None
    for cat,mod,desc,how,exp in FIXES:
        if cat!=cur:
            h.append(f"<tr><td colspan='4' style='background:#fdf3e3;color:#b45309;font-weight:700;border-top:2px solid #b45309'>{cat}</td></tr>")
            cur=cat
        h.append(f"<tr><td><b>{mod}</b></td><td>{desc}</td><td>{how}</td><td>{exp}</td></tr>")
    h.append("</table></div>")

    h.append(f"<footer>เอกสารสร้างอัตโนมัติจากประวัติการพัฒนา Herbal ERP • ช่วง {PERIOD}</footer>")
    h.append("</div></body></html>")
    path = os.path.join(OUT,"system-flow_HerbalERP.html")
    with open(path,"w",encoding="utf-8") as fp: fp.write("\n".join(h))
    saved("HTML (System Flow)", path)

# ====================================================================
# 2) DOCX — คู่มือการบันทึกข้อมูล_HerbalERP.docx
# ====================================================================
def build_docx():
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    FONT="TH Sarabun New"
    doc = Document()
    # default style
    st = doc.styles["Normal"]; st.font.name=FONT; st.font.size=Pt(15)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),FONT); st.element.rPr.rFonts.set(qn('w:cs'),FONT)

    def setfont(run, sz=15, bold=False, color=None):
        run.font.name=FONT; run.font.size=Pt(sz); run.font.bold=bold
        if color: run.font.color.rgb=RGBColor(*color)
        rpr=run._element.get_or_add_rPr()
        rf=rpr.find(qn('w:rFonts'))
        if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),FONT)

    def para(text="", sz=15, bold=False, color=None, align=None, space=4):
        p=doc.add_paragraph();
        if align: p.alignment=align
        p.paragraph_format.space_after=Pt(space)
        r=p.add_run(text); setfont(r,sz,bold,color); return p

    def bullet(text, sz=15):
        p=doc.add_paragraph(style="List Bullet"); r=p.add_run(text); setfont(r,sz);
        p.paragraph_format.space_after=Pt(2); return p

    NAVY=(0x1e,0x3a,0x5f); TEAL=(0x0f,0x76,0x6e); AMBER=(0xb4,0x53,0x09)

    # cover
    para("คู่มือการบันทึกข้อมูล", 34, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 2)
    para("ระบบ Herbal ERP", 26, True, TEAL, WD_ALIGN_PARAGRAPH.CENTER, 2)
    para("สำหรับผู้ใช้งาน — แนวทางการบันทึกข้อมูลแต่ละโมดูล", 16, False, None, WD_ALIGN_PARAGRAPH.CENTER, 2)
    para("ครอบคลุมระบบที่พัฒนาในช่วง "+PERIOD, 15, False, (0x47,0x55,0x69), WD_ALIGN_PARAGRAPH.CENTER, 14)
    try: doc.add_picture(IMGS["overview"], width=Inches(6.3))
    except Exception: pass
    doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # how to read
    para("วิธีใช้คู่มือเล่มนี้", 20, True, NAVY)
    bullet("แต่ละหัวข้อคือ 1 โมดูล ระบุ 'เมนู' (เส้นทางคลิก) วัตถุประสงค์ และขั้นตอนการบันทึกข้อมูล")
    bullet("ขั้นตอนแบ่งเป็น เพิ่ม (สร้างใหม่) / แก้ไข (ปรับข้อมูลเดิม) / ลบ (นำออก)")
    bullet("รูปประกอบแสดงลำดับการทำงาน (flow) ของโมดูลนั้น ๆ")
    doc.add_page_break()

    for i,m in enumerate(MODULES,1):
        para(f"{i}. {m['name']}", 20, True, NAVY, space=2)
        nav_p = doc.add_paragraph(); nav_p.paragraph_format.space_after=Pt(1)
        r=nav_p.add_run("วิธีเข้าถึงหน้าจอ:  "); setfont(r,15,True,NAVY)
        r=nav_p.add_run(f"เมนูซ้าย ▸ {m['module']} ▸ {m['menu']}"); setfont(r,15,False,TEAL)
        para(f"URL อ้างอิง:  {m['url']}", 13, False, (0x47,0x55,0x69), space=2)
        if m.get("nav2"):
            para(f"เพิ่มเติม:  {m['nav2']}", 13, False, (0x47,0x55,0x69), space=2)
        para(f"วัตถุประสงค์:  {m['purpose']}", 15, space=6)
        if m.get("diagram"):
            try:
                doc.add_picture(IMGS[m["diagram"]], width=Inches(6.0))
                doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
            except Exception: pass
        for label,key,col in [("การเพิ่มข้อมูล (Create)","add",TEAL),
                              ("การแก้ไขข้อมูล (Update)","edit",AMBER),
                              ("การลบข้อมูล (Delete)","delete",(0xb9,0x1c,0x1c))]:
            if m.get(key):
                para(label, 16, True, col, space=2)
                for s in m[key]: bullet(s)
        if i < len(MODULES): doc.add_page_break()

    doc.add_page_break()
    para("ภาคผนวก — สรุประบบที่ปรับแก้ไข", 20, True, NAVY)
    para("รายการต่อไปนี้คือการปรับปรุง/แก้ไขระบบเดิมในช่วงเดียวกัน", 15, space=8)
    table = doc.add_table(rows=1, cols=4); table.style="Light Grid Accent 1"
    hdr=table.rows[0].cells
    for c,t in zip(hdr,["หมวด","โมดูล","สิ่งที่ปรับแก้ไข","ผลที่คาดหวัง"]):
        c.paragraphs[0].clear(); r=c.paragraphs[0].add_run(t); setfont(r,14,True,(0xff,0xff,0xff))
    for cat,mod,desc,how,exp in FIXES:
        row=table.add_row().cells
        for c,t in zip(row,[cat,mod,desc,exp]):
            c.paragraphs[0].clear(); r=c.paragraphs[0].add_run(t); setfont(r,13)
    path=os.path.join(OUT,"คู่มือการบันทึกข้อมูล_HerbalERP.docx")
    doc.save(path); saved("DOCX (คู่มือการบันทึกข้อมูล)", path)

# ====================================================================
# 3) XLSX — แบบทดสอบระบบ_UAT_HerbalERP.xlsx
# ====================================================================
def build_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    FONT="TH Sarabun New"
    wb=Workbook()
    navy=PatternFill("solid",fgColor="1E3A5F"); teal=PatternFill("solid",fgColor="0F766E")
    amber=PatternFill("solid",fgColor="B45309"); grey=PatternFill("solid",fgColor="F1F5F9")
    band=PatternFill("solid",fgColor="F7FAFF")
    thin=Side(style="thin",color="C8D3E0"); border=Border(thin,thin,thin,thin)
    wrap=Alignment(wrap_text=True,vertical="top")
    hfont=Font(name=FONT,size=14,bold=True,color="FFFFFF")
    tfont=Font(name=FONT,size=14,bold=True,color="1E3A5F")
    bfont=Font(name=FONT,size=13)

    def style_header(ws,row,fill):
        for c in ws[row]:
            c.fill=fill; c.font=hfont; c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center"); c.border=border

    # ---- cover sheet ----
    ws0=wb.active; ws0.title="ปก"
    ws0.sheet_view.showGridLines=False
    ws0["B2"]="แบบทดสอบระบบ (UAT)"; ws0["B2"].font=Font(name=FONT,size=26,bold=True,color="1E3A5F")
    ws0["B3"]="ระบบ Herbal ERP"; ws0["B3"].font=Font(name=FONT,size=20,bold=True,color="0F766E")
    ws0["B4"]=f"ช่วงการพัฒนา: {PERIOD}"; ws0["B4"].font=Font(name=FONT,size=14)
    info=["","วิธีใช้แบบทดสอบ:",
          "1) ไปที่ชีต 'ระบบพัฒนาเพิ่ม-CRUD' เพื่อทดสอบฟังก์ชัน เพิ่ม / แก้ไข / ลบ",
          "2) ไปที่ชีต 'ปรับแก้ไขระบบ' เพื่อตรวจการแก้ไขระบบเดิม",
          "3) ทำตามคอลัมน์ 'ขั้นตอนการคลิก' แล้วเทียบกับ 'ผลที่คาดหวัง'",
          "4) เลือกผลในคอลัมน์ 'ผลทดสอบ' = Pass / Fail / N/A (มีดรอปดาวน์ให้เลือก)",
          "5) กรอกผู้ทดสอบ / วันที่ / หมายเหตุ เมื่อพบปัญหา",
          "",
          "หมายเหตุ: เส้นทางเมนู (เช่น /quality/qc-entry) คือ URL ที่พิมพ์ต่อท้ายที่อยู่เว็บ หรือคลิกผ่านเมนูซ้าย"]
    r=6
    for line in info:
        ws0[f"B{r}"]=line; ws0[f"B{r}"].font=Font(name=FONT,size=14,bold=(line.endswith(":"))); r+=1
    ws0.column_dimensions["A"].width=3; ws0.column_dimensions["B"].width=110

    # ---- sheet CRUD ----
    ws=wb.create_sheet("ระบบพัฒนาเพิ่ม-CRUD")
    cols=["ลำดับ","รหัส","Module (เมนูหลัก)","เมนูที่คลิก","วิธีเข้าถึงหน้าจอ (คลิกตามนี้)","URL อ้างอิง",
          "ฟังก์ชัน","ขั้นตอนการทดสอบ","ผลที่คาดหวัง","ผลทดสอบ","ผู้ทดสอบ","วันที่","หมายเหตุ"]
    ws.append(cols); style_header(ws,1,navy)
    widths=[6,7,26,26,40,30,11,46,44,11,14,12,26]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A2"
    fillmap={"เพิ่ม":"E8F5EE","แก้ไข":"FDF3E3","ลบ":"FDECEC","ดู":"EEF2F7"}
    n=1; rr=2
    for m in MODULES:
        mod_short=m["module"].split("(")[0].strip()
        access=f"เมนูซ้าย ▸ {mod_short} ▸ {m['menu']}"
        first=True
        for fn,step,exp in m["tests"]:
            # ใส่ข้อมูล Module/เมนู เฉพาะแถวแรกของแต่ละโมดูล เพื่อให้อ่านง่ายเป็นบล็อก
            ws.append([n, m["code"] if first else "", m["module"] if first else "",
                       m["menu"] if first else "", access if first else "", m["url"] if first else "",
                       fn, step, exp, "", "", "", ""])
            for c in ws[rr]:
                c.font=bfont; c.alignment=wrap; c.border=border
            ws[f"G{rr}"].fill=PatternFill("solid",fgColor=fillmap.get(fn,"FFFFFF"))
            ws[f"G{rr}"].alignment=Alignment(horizontal="center",vertical="center")
            ws[f"J{rr}"].alignment=Alignment(horizontal="center",vertical="center")
            if first:
                for col in ("C","D","E","F"):
                    ws[f"{col}{rr}"].font=Font(name=FONT,size=13,bold=True,color="1E3A5F")
            n+=1; rr+=1; first=False
    dv=DataValidation(type="list",formula1='"Pass,Fail,N/A"',allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"J2:J{rr-1}")

    # ---- sheet fixes ----
    ws2=wb.create_sheet("ปรับแก้ไขระบบ")
    cols2=["ลำดับ","หมวด","โมดูล","สิ่งที่ปรับแก้ไข","วิธีตรวจสอบ (คลิกตามนี้)","ผลที่คาดหวัง","ผลทดสอบ","ผู้ทดสอบ","วันที่","หมายเหตุ"]
    ws2.append(cols2); style_header(ws2,1,amber)
    w2=[6,30,24,50,46,40,11,14,12,26]
    for i,w in enumerate(w2,1): ws2.column_dimensions[chr(64+i)].width=w
    ws2.freeze_panes="A2"
    rr=2; cur=None
    for i,(cat,mod,desc,how,exp) in enumerate(FIXES,1):
        ws2.append([i, cat if cat!=cur else "", mod, desc, how, exp, "","","",""])
        for c in ws2[rr]: c.font=bfont; c.alignment=wrap; c.border=border
        if cat!=cur:
            ws2[f"B{rr}"].font=Font(name=FONT,size=13,bold=True,color="B45309"); cur=cat
        ws2[f"G{rr}"].alignment=Alignment(horizontal="center",vertical="center")
        rr+=1
    dv2=DataValidation(type="list",formula1='"Pass,Fail,N/A"',allow_blank=True)
    ws2.add_data_validation(dv2); dv2.add(f"G2:G{rr-1}")

    path=os.path.join(OUT,"แบบทดสอบระบบ_UAT_HerbalERP.xlsx")
    wb.save(path); saved("XLSX (แบบทดสอบ UAT)", path)

# ====================================================================
# 4) PPTX — สรุประบบ_HerbalERP.pptx
# ====================================================================
def build_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.ns import qn
    FONT="TH Sarabun PSK"
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    SW,SH=prs.slide_width,prs.slide_height
    blank=prs.slide_layouts[6]
    NAVY=RGBColor(0x1e,0x3a,0x5f); BLUE=RGBColor(0x2c,0x6f,0xbf); TEAL=RGBColor(0x0f,0x76,0x6e)
    WHITE=RGBColor(0xff,0xff,0xff); SLATE=RGBColor(0x47,0x55,0x69); AMBER=RGBColor(0xb4,0x53,0x09)
    GREEN=RGBColor(0x15,0x80,0x3d)

    def thai(run, sz, bold=False, color=None):
        run.font.size=Pt(sz); run.font.bold=bold; run.font.name=FONT
        if color is not None: run.font.color.rgb=color
        rPr=run._r.get_or_add_rPr()
        for tag in ("a:latin","a:ea","a:cs"):
            el=rPr.find(qn(tag))
            if el is None:
                el=rPr.makeelement(qn(tag),{}); rPr.append(el)
            el.set("typeface",FONT)

    def textbox(slide,l,t,w,h,lines,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
        tb=slide.shapes.add_textbox(l,t,w,h); tf=tb.text_frame; tf.word_wrap=True
        tf.vertical_anchor=anchor
        for i,(txt,sz,bold,col) in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.alignment=align; r=p.add_run(); r.text=txt; thai(r,sz,bold,col)
        return tb

    def bg(slide,color):
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color

    def bar(slide,title,color=NAVY):
        sh=slide.shapes.add_shape(1,0,0,SW,Inches(1.05))
        sh.fill.solid(); sh.fill.fore_color.rgb=color; sh.line.fill.background()
        tf=sh.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.4)
        tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        r=tf.paragraphs[0].add_run(); r.text=title; thai(r,30,True,WHITE)

    def pic_fit(slide, key, top, max_w, max_h, cx=None):
        from PIL import Image
        iw,ih=Image.open(IMGS[key]).size
        ratio=min(max_w/iw, max_h/ih); w=int(iw*ratio); h=int(ih*ratio)
        left = (SW-w)//2 if cx is None else cx
        slide.shapes.add_picture(IMGS[key], left, top, width=w, height=h)

    # slide 1 title
    s=prs.slides.add_slide(blank); bg(s,NAVY)
    box=s.shapes.add_shape(1,Inches(0),Inches(2.4),SW,Inches(2.7))
    box.fill.solid(); box.fill.fore_color.rgb=NAVY; box.line.fill.background()
    textbox(s,Inches(0.6),Inches(2.5),Inches(12.1),Inches(2.6),[
        ("สรุประบบ Herbal ERP",46,True,WHITE),
        ("ระบบบริหารการผลิตยาสมุนไพรตามมาตรฐาน GMP",26,False,RGBColor(0xcf,0xe2,0xf7)),
        ("สรุปการพัฒนาในช่วง "+PERIOD,20,False,RGBColor(0xa9,0xc7,0xe8)),
    ],PP_ALIGN.CENTER)

    # slide 2 overview
    s=prs.slides.add_slide(blank); bg(s,WHITE); bar(s,"ภาพรวมระบบ — วงจรการผลิตตาม GMP")
    pic_fit(s,"overview",Inches(1.3),int(Inches(12.3)),int(Inches(5.6)))

    # slide 3 list of new modules
    s=prs.slides.add_slide(blank); bg(s,WHITE); bar(s,"ระบบที่พัฒนาเพิ่ม (12 ระบบ)",TEAL)
    lines=[]
    for i,m in enumerate(MODULES,1):
        nm=m["name"].split("(")[0].strip()
        lines.append((f"{i}.  {nm}   {m['code'] if m['code']!='—' else ''}",17,False,NAVY))
    half=(len(lines)+1)//2
    textbox(s,Inches(0.6),Inches(1.4),Inches(6.2),Inches(5.6),lines[:half],PP_ALIGN.LEFT)
    textbox(s,Inches(6.9),Inches(1.4),Inches(6.2),Inches(5.6),lines[half:],PP_ALIGN.LEFT)

    # per-module slides that have diagrams
    diagram_modules=[m for m in MODULES if m.get("diagram")]
    for m in diagram_modules:
        s=prs.slides.add_slide(blank); bg(s,WHITE)
        nm=m["name"]
        bar(s, nm if len(nm)<48 else nm[:46]+"…", BLUE)
        textbox(s,Inches(0.6),Inches(1.15),Inches(12.1),Inches(0.9),
                [(m["purpose"],16,False,SLATE)],PP_ALIGN.LEFT)
        pic_fit(s,m["diagram"],Inches(2.1),int(Inches(12.0)),int(Inches(4.7)))

    # fixes summary slide
    s=prs.slides.add_slide(blank); bg(s,WHITE); bar(s,"ระบบที่ปรับแก้ไข (สรุปตามหมวด)",AMBER)
    flines=[]; cur=None
    for cat,mod,desc,how,exp in FIXES:
        if cat!=cur:
            flines.append((cat,16,True,AMBER)); cur=cat
        flines.append((f"   •  {mod}: {desc}",12,False,NAVY))
    half=(len(flines)+1)//2
    # หาจุดตัดที่ไม่ผ่ากลางหัวข้อหมวด
    while half<len(flines) and flines[half][2] is False: half+=1
    textbox(s,Inches(0.5),Inches(1.3),Inches(6.3),Inches(5.9),flines[:half],PP_ALIGN.LEFT)
    textbox(s,Inches(6.9),Inches(1.3),Inches(6.0),Inches(5.9),flines[half:],PP_ALIGN.LEFT)

    # closing
    s=prs.slides.add_slide(blank); bg(s,NAVY)
    textbox(s,Inches(0.8),Inches(2.6),Inches(11.7),Inches(2.4),[
        ("ขั้นตอนถัดไป: การทดสอบระบบ (UAT)",34,True,WHITE),
        ("ใช้ไฟล์ 'แบบทดสอบระบบ_UAT_HerbalERP.xlsx' ทดสอบฟังก์ชัน เพิ่ม / แก้ไข / ลบ ของแต่ละโมดูล",20,False,RGBColor(0xcf,0xe2,0xf7)),
        ("และตรวจรายการปรับแก้ไขระบบตามชีต 'ปรับแก้ไขระบบ'",20,False,RGBColor(0xcf,0xe2,0xf7)),
    ],PP_ALIGN.CENTER)

    path=os.path.join(OUT,"สรุประบบ_HerbalERP.pptx")
    prs.save(path); saved("PPTX (สไลด์นำเสนอ)", path)

# ---- run all with per-file error handling ----
for name,fn in [("HTML",build_html),("DOCX",build_docx),("XLSX",build_xlsx),("PPTX",build_pptx)]:
    try:
        fn()
        print(f"[OK]   {name}")
    except PermissionError as e:
        print(f"[LOCKED] {name}: ไฟล์เปิดอยู่ใน Office — {e}")
    except Exception as e:
        import traceback; print(f"[FAIL] {name}: {e}"); traceback.print_exc()

print("\n===== RESULT =====")
for nm,p,sz in results:
    print(f"  {nm}: {p}  ({sz//1024} KB)")
